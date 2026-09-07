from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook


@dataclass
class ImportReview:
    sheet: str
    header_row: int
    frame: pd.DataFrame
    excluded: list[dict[str, Any]] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    mapping: dict[str, str] = field(default_factory=dict)


def _names(values: list[Any]) -> list[str]:
    seen: dict[str, int] = {}
    result = []
    for i, value in enumerate(values):
        base = str(value).strip() if value not in (None, "") else f"coluna_{i + 1}"
        n = seen.get(base, 0)
        seen[base] = n + 1
        result.append(base if n == 0 else f"{base}__{n + 1}")
    return result


def inspect_workbook(path: str | Path) -> dict[str, Any]:
    try:
        wb_formula = load_workbook(path, read_only=True, data_only=False, keep_vba=False)
        wb_values = load_workbook(path, read_only=True, data_only=True, keep_vba=False)
    except (PermissionError, OSError) as exc:
        raise ValueError("Arquivo indisponível localmente. No OneDrive, mantenha-o neste dispositivo.") from exc
    sheets, formula_without_cache = [], []
    for wf, wv in zip(wb_formula.worksheets, wb_values.worksheets):
        sheets.append({"name": wf.title, "rows": wf.max_row, "columns": wf.max_column})
        for rf, rv in zip(wf.iter_rows(), wv.iter_rows()):
            for cf, cv in zip(rf, rv):
                if isinstance(cf.value, str) and cf.value.startswith("=") and cv.value is None:
                    formula_without_cache.append(f"{wf.title}!{cf.coordinate}")
    return {"sheets": sheets, "formula_without_cache": formula_without_cache[:100]}


def detect_header(path: str | Path, sheet: str, max_rows: int = 40) -> int:
    raw = pd.read_excel(path, sheet_name=sheet, header=None, nrows=max_rows, dtype=object)
    tokens = ("data", "hora", "lote", "ordem", "gramatura", "receita", "parâmetro")
    scores = [sum(any(t in str(v).lower() for t in tokens) for v in row if pd.notna(v)) for row in raw.values]
    return int(max(range(len(scores)), key=scores.__getitem__))


def import_sheet(path: str | Path, sheet: str, header_row: int | None = None) -> ImportReview:
    header_row = detect_header(path, sheet) if header_row is None else header_row
    raw = pd.read_excel(path, sheet_name=sheet, header=None, dtype=object)
    names = _names(raw.iloc[header_row].tolist())
    data = raw.iloc[header_row + 1 :].copy()
    data.columns = names
    data.insert(0, "_source_row", range(header_row + 2, header_row + 2 + len(data)))
    excluded, keep = [], []
    header_norm = {re.sub(r"\s+", " ", n.strip().lower()) for n in names}
    for idx, row in data.iterrows():
        populated = [v for k, v in row.items() if k != "_source_row" and pd.notna(v) and str(v).strip()]
        vals = {re.sub(r"\s+", " ", str(v).strip().lower()) for v in populated}
        reason = None
        if not populated:
            reason = "linha vazia"
        elif len(vals & header_norm) >= max(2, min(4, len(populated) // 2)):
            reason = "cabeçalho repetido"
        elif any(str(v).strip().lower() in {"média", "media", "desvio padrão", "target", "mínimo", "máximo"} for v in populated[:3]):
            reason = "estatística agregada"
        if reason:
            excluded.append({"row": int(row["_source_row"]), "reason": reason})
        else:
            keep.append(idx)
    frame = data.loc[keep].reset_index(drop=True)
    warnings = []
    duplicated = int(frame.drop(columns="_source_row").duplicated().sum())
    if duplicated:
        warnings.append(f"{duplicated} linhas duplicadas preservadas para revisão")
    return ImportReview(sheet, header_row + 1, frame, excluded, warnings)
