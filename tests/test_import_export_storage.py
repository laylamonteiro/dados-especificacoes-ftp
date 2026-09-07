import hashlib
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook

from receita_local.exports.writers import excel_bytes, pdf_bytes
from receita_local.importers.excel import import_sheet, inspect_workbook
from receita_local.storage.repository import Repository


def fixture_xlsx(path: Path):
    wb=Workbook(); ws=wb.active; ws.title="PARAMETROS DE PROCESSO "
    ws.append(["regra introdutória"]); ws.append([None]); ws.append(["Data","Valor","Valor"])
    ws.append(["2026-01-01",1,None]); ws.append(["Data","Valor","Valor"]); ws.append(["Média",1,2]); ws.append(["2026-01-02",100,"=1+1"])
    wb.save(path)


def test_shifted_repeated_duplicate_headers_stats_formula_cache(tmp_path):
    path=tmp_path/"caminho com acento ç.xlsx"; fixture_xlsx(path)
    review=import_sheet(path,"PARAMETROS DE PROCESSO ")
    assert review.header_row == 3 and list(review.frame.columns)[1:] == ["Data","Valor","Valor__2"]
    assert len(review.frame)==2 and {x["reason"] for x in review.excluded} == {"cabeçalho repetido","estatística agregada"}
    assert inspect_workbook(path)["formula_without_cache"] == ["PARAMETROS DE PROCESSO !C7"]


def test_storage_preserves_original_deduplicates_persists_and_exports(tmp_path):
    source=tmp_path/"original.xlsx"; fixture_xlsx(source); before=hashlib.sha256(source.read_bytes()).hexdigest()
    repo=Repository(tmp_path/"dados"); one=repo.import_file(source); two=repo.import_file(source)
    assert not one["duplicate"] and two["duplicate"] and before==hashlib.sha256(source.read_bytes()).hexdigest()
    result={"mode":"exploratório","items":[{"machine":"M","parameter":"P","unit":"C","target":2.0,"lower":1.0,"upper":3.0,"status":"proposto"}],"evidence":["x"],"assumptions":[],"pending":[],"quality":[],"algorithm_version":"v"}
    aid=repo.save_analysis("001",{},result,[one["id"]]); repo.db.close(); reopened=Repository(tmp_path/"dados")
    assert reopened.get_analysis(aid)["items"] == result["items"]
    excel=excel_bytes(result,{"id":aid}); pdf=pdf_bytes(result,{"id":aid})
    assert excel.startswith(b"PK") and pdf.startswith(b"%PDF")
    out=tmp_path/"out.xlsx"; out.write_bytes(excel); assert load_workbook(out)["Receita"]["D2"].value == 2
    reopened.delete_analysis(aid); assert not Path(one["stored_path"]).exists()
